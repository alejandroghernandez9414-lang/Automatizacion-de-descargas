"""
Descarga los 3 Excel de ordenes (DC00, RC00, ZVCL) desde la plataforma web
de gestion de trabajo en campo (WTM),
luego genera BASES_PSR.xlsx con las hojas BD-WTM y BD-SESAP aplicando
la misma logica de las macros VBA del KPI (orden, OrdRep_WTM,
EstadoTrabajo_WTM, cruces con SESAP).

USO DIARIO (mañana - descargar + generar):
    python descargar_y_generar_bases.py

Cuando solo cambia el SESAP (sin re-descargar WTM):
    python generar_bases.py

Requisitos:
    pip install -r requirements.txt
    playwright install chromium
    Copiar .env.example como .env y completar los valores.
"""

from playwright.sync_api import sync_playwright
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
from openpyxl.utils import get_column_letter
import glob
import os
import unicodedata
import pandas as pd
import openpyxl

from config import requerido

# ------------------------------------------------------------------
# CONFIGURACION
# ------------------------------------------------------------------
# Todo lo sensible (usuario, clave, rutas e identificadores) vive en .env
USUARIO = requerido("WTM_USUARIO")
CLAVE   = requerido("WTM_CLAVE")

CARPETA_BD_WTM = requerido("CARPETA_BD_WTM")
CARPETA_SESAP  = requerido("CARPETA_SESAP")
RUTA_SALIDA    = requerido("RUTA_SALIDA")

FILTRO_ESTADO_PROCESOFINAL = "ODS ASIGNADA"   # filtro sobre BD_SESAP_GBR

WTM_BASE_URL    = requerido("WTM_BASE_URL").rstrip("/")
INDEX_URL       = f"{WTM_BASE_URL}/FwAccount/Index"
MYPROCESSES_URL = f"{WTM_BASE_URL}/forminstances/myprocesses"

PROCESOS = {
    "DC00": {
        "nombre_largo": "Suspensión por deuda DC00",
        "report_url":   f"{MYPROCESSES_URL}/Report?id={requerido('WTM_REPORT_ID_DC00')}",
        "export_data_id": requerido("WTM_EXPORT_ID_DC00"),
        "patron_archivo_viejo": "*Suspensión por deuda DC00*.xlsx",
    },
    "RC00": {
        "nombre_largo": "Reconexión por pago RC00",
        "report_url":   f"{MYPROCESSES_URL}/Report?id={requerido('WTM_REPORT_ID_RC00')}",
        "export_data_id": requerido("WTM_EXPORT_ID_RC00"),
        "patron_archivo_viejo": "*Reconexión por pago RC00*.xlsx",
    },
    "ZVCL": {
        "nombre_largo": "Visita a Cliente ZVCL",
        "report_url":   f"{MYPROCESSES_URL}/Report?id={requerido('WTM_REPORT_ID_ZVCL')}",
        "export_data_id": requerido("WTM_EXPORT_ID_ZVCL"),
        "patron_archivo_viejo": "*Visita a Cliente ZVCL*.xlsx",
    },
}

# Columnas que se incluyen en cada hoja del BASES_PSR.xlsx (mismo orden que el KPI)
COLUMNAS_WTM = [
    "IdentificadorM","NumeroOrdenWTM","TipoOrden","NombreFormulario","EstadoFormulario",
    "FormularioCompletado","NombreAgente","CodigoAgente","FechaAsignacion",
    "FechaLimite","FechaInicio","FechaCompletado","TrabajoCompletado",
    "Latitud","Longitud","NIT","Nombre","Direccion","Barrio","Ciudad",
    "Departamento","TipoDocumento","NombreFormularioSiguiente","NumeroFormulariosSiguientes",
    "Grupo","OrdRep_WTM","EstadoTrabajo_WTM","F_TRABAJO_SESAP",
]

COLUMNAS_SESAP = [
    "NumeroOrden","ClaseDeOrden","Cod_TipoIdent","TipoIdentificacion",
    "NumeroIdentificacion","NombreCliente","DireccionInstalacion","Municipio",
    "Estrato","Ciclo","FechaCreacion","FechaModificacion","UsuarioCreacion",
    "EstadoProceso","ESTADO_PROCESOFINAL_GBR","F_TRABAJO_PLAN",
    "EstadoTrabajoFin_WTM","Estado_FHMAX","CT_FHMAX","FHMAX",
]


# ------------------------------------------------------------------
# DESCARGA (Playwright)
# ------------------------------------------------------------------

def borrar_descargas_viejas(patron: str):
    ruta_patron = os.path.join(CARPETA_BD_WTM, patron)
    for viejo in glob.glob(ruta_patron):
        os.remove(viejo)
        print(f"  Borrado archivo anterior: {os.path.basename(viejo)}")


def descargar_proceso(clave: str, datos: dict):
    def log(msg):
        print(f"[{clave}] {msg}")

    log(f"=== {datos['nombre_largo']} ===")

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(accept_downloads=True)
        page = context.new_page()

        # 1. Login
        log("Iniciando sesion...")
        page.goto(INDEX_URL, timeout=60000)
        page.fill("#UserName", USUARIO)
        page.fill("#Password", CLAVE)
        page.click("#loginForm button[data-action='ajax-send']")
        page.wait_for_url("**/Home/Index", timeout=60000)

        # 2. Navegar al reporte
        page.goto(MYPROCESSES_URL, timeout=60000)
        page.goto(datos["report_url"], timeout=120000)
        page.wait_for_selector("[id$='-gridrangefilter']", state="visible", timeout=60000)

        # 3. Filtrar: clic en "Últimos 30 días" (cierra solo, sin Aplicar)
        log("Filtrando fechas: Últimos 30 días")
        page.click("[id$='-gridrangefilter']")
        page.wait_for_selector(".daterangepicker", state="visible")
        panel = page.locator(".daterangepicker:visible")
        panel.locator("text=Últimos 30 días").click()
        page.wait_for_selector(".daterangepicker", state="hidden", timeout=15000)

        # 4. Exportar
        log("Exportando (puede tardar varios minutos)...")
        page.click("#tableTools-export-types")
        page.wait_for_selector(".dropdown-menu.open, .open .dropdown-menu", state="visible")
        with page.expect_download(timeout=3_000_000) as download_info:
            page.click(f"a[data-id='{datos['export_data_id']}']")

        download = download_info.value
        nombre_archivo = download.suggested_filename
        ruta_destino = os.path.join(CARPETA_BD_WTM, nombre_archivo)

        # 5. Borrar archivo anterior y guardar el nuevo
        borrar_descargas_viejas(datos["patron_archivo_viejo"])
        download.save_as(ruta_destino)
        log(f"Guardado: {ruta_destino}")

        browser.close()


# ------------------------------------------------------------------
# GENERACION DE BASES (pandas + openpyxl)
# ------------------------------------------------------------------

def _normalizar(texto: str) -> str:
    """Quita tildes y pone en minusculas para comparar columnas."""
    nfkd = unicodedata.normalize("NFKD", str(texto))
    sin_tildes = "".join(c for c in nfkd if not unicodedata.combining(c))
    return sin_tildes.lower().strip()


def _leer_hoja(ruta: str, columnas_deseadas: list[str]) -> pd.DataFrame:
    """Lee solo las columnas necesarias de un Excel usando openpyxl read_only.
    Tolera diferencias de mayusculas/minusculas y tildes en los encabezados."""
    norm_deseadas = {_normalizar(c): c for c in columnas_deseadas}

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb.active

    filas = ws.iter_rows(values_only=True)
    encabezados_raw = next(filas, None)
    if encabezados_raw is None:
        wb.close()
        return pd.DataFrame(columns=columnas_deseadas)

    # Mapear indice -> nombre canonico de columna
    idx_a_col = {}
    for i, h in enumerate(encabezados_raw):
        if h is None:
            continue
        norm = _normalizar(str(h))
        if norm in norm_deseadas:
            idx_a_col[i] = norm_deseadas[norm]

    registros = []
    for fila in filas:
        registro = {col: None for col in columnas_deseadas}
        for idx, col in idx_a_col.items():
            if idx < len(fila):
                registro[col] = fila[idx]
        registros.append(registro)

    wb.close()
    return pd.DataFrame(registros, columns=columnas_deseadas)


def construir_wtm() -> pd.DataFrame:
    """Lee y concatena los archivos DC00, RC00 y ZVCL de CARPETA_BD_WTM."""
    archivos = glob.glob(os.path.join(CARPETA_BD_WTM, "*.xlsx"))
    if not archivos:
        raise FileNotFoundError(f"No se encontraron archivos .xlsx en {CARPETA_BD_WTM}")

    columnas_base = [
        "IdentificadorM","NumeroOrdenWTM","TipoOrden","NombreFormulario","EstadoFormulario",
        "FormularioCompletado","NombreAgente","CodigoAgente","FechaAsignacion",
        "FechaLimite","FechaInicio","FechaCompletado","TrabajoCompletado",
        "Latitud","Longitud","NIT","Nombre","Direccion","Barrio","Ciudad",
        "Departamento","TipoDocumento","NombreFormularioSiguiente","NumeroFormulariosSiguientes",
        "Grupo",
    ]

    partes = []
    for ruta in archivos:
        df = _leer_hoja(ruta, columnas_base)

        # Convertir TrabajoCompletado: Sí/Si -> True, No -> False
        def conv_tc(v):
            if v is None:
                return None
            s = str(v).strip().lower()
            if s in ("sí", "si", "true", "verdadero", "1"):
                return True
            if s in ("no", "false", "falso", "0"):
                return False
            return v

        df["TrabajoCompletado"] = df["TrabajoCompletado"].apply(conv_tc)
        partes.append(df)

    wtm = pd.concat(partes, ignore_index=True)

    # --- Ordenar: NumeroOrdenWTM ASC, IdentificadorM DESC (igual que OrdenarTbl_BD_WTM) ---
    wtm = wtm.sort_values(
        ["NumeroOrdenWTM", "IdentificadorM"],
        ascending=[True, False],
        na_position="last",
    ).reset_index(drop=True)

    # --- OrdRep_WTM ---
    def calc_ord_rep(row, prev_orden, prev_id):
        orden_actual = row["NumeroOrdenWTM"]
        id_actual    = row["IdentificadorM"]
        if pd.notna(orden_actual) and pd.notna(prev_orden) and orden_actual == prev_orden:
            if pd.notna(id_actual) and pd.notna(prev_id) and id_actual < prev_id:
                return "OrdRepetida_IMM"
        return ""

    ord_rep = []
    prev_orden = prev_id = None
    for _, row in wtm.iterrows():
        val = calc_ord_rep(row, prev_orden, prev_id)
        ord_rep.append(val)
        prev_orden = row["NumeroOrdenWTM"]
        prev_id    = row["IdentificadorM"]
    wtm["OrdRep_WTM"] = ord_rep

    # --- EstadoTrabajo_WTM ---
    def calc_estado(row):
        tc = row["TrabajoCompletado"]
        fc = str(row["FormularioCompletado"]).strip() if pd.notna(row["FormularioCompletado"]) else ""
        if tc is True  and fc == "Cancelado": return "CANCELADO"
        if tc is False and fc == "Cancelado": return "CANCELADO"
        if tc is True  and fc == "Cerrado":   return "EJECUTADO"
        if tc is False and fc == "Cerrado":   return "PENDIENTE"
        return "."

    wtm["EstadoTrabajo_WTM"] = wtm.apply(calc_estado, axis=1)

    # F_TRABAJO_SESAP se llena en aplicar_cruces()
    wtm["F_TRABAJO_SESAP"] = pd.NaT

    print(f"Total BD-WTM: {len(wtm)} filas ({len(archivos)} archivos)")
    return wtm


def construir_sesap() -> pd.DataFrame:
    """Lee el archivo SESAP mas reciente de CARPETA_SESAP."""
    archivos = sorted(glob.glob(os.path.join(CARPETA_SESAP, "*.xlsx")))
    if not archivos:
        raise FileNotFoundError(f"No se encontraron archivos .xlsx en {CARPETA_SESAP}")

    ruta = archivos[-1]   # el mas reciente (por nombre)
    columnas_base = [
        "NumeroOrden","ClaseDeOrden","Cod_TipoIdent","TipoIdentificacion",
        "NumeroIdentificacion","NombreCliente","DireccionInstalacion","Municipio",
        "Estrato","Ciclo","FechaCreacion","FechaModificacion","UsuarioCreacion",
        "EstadoProceso","ESTADO_PROCESOFINAL_GBR","F_TRABAJO_PLAN","FHMAX",
    ]

    sesap = _leer_hoja(ruta, columnas_base)

    # Filtrar por estado
    if FILTRO_ESTADO_PROCESOFINAL:
        antes = len(sesap)
        sesap = sesap[
            sesap["ESTADO_PROCESOFINAL_GBR"].astype(str).str.strip() == FILTRO_ESTADO_PROCESOFINAL
        ].reset_index(drop=True)
        print(f"  SESAP filtrado '{FILTRO_ESTADO_PROCESOFINAL}': {antes} -> {len(sesap)} filas")

    # Columnas calculadas; se llenan en aplicar_cruces()
    sesap["EstadoTrabajoFin_WTM"] = ""
    sesap["Estado_FHMAX"]         = ""
    sesap["CT_FHMAX"]             = ""

    print(f"Total BD-SESAP: {len(sesap)} filas ({os.path.basename(ruta)})")
    return sesap


def aplicar_cruces(wtm: pd.DataFrame, sesap: pd.DataFrame):
    """Replica la logica de AplicarFuncionesBDWTM y AplicarFuncionesBDSESAP."""
    import datetime

    # --- F_TRABAJO_SESAP (col 37 de BD_SESAP_GBR = F_TRABAJO_PLAN) ---
    tabla_sesap_f = (
        sesap.drop_duplicates(subset=["NumeroOrden"], keep="first")
             .set_index("NumeroOrden")["F_TRABAJO_PLAN"]
    )
    wtm["F_TRABAJO_SESAP"] = wtm["NumeroOrdenWTM"].map(tabla_sesap_f)
    wtm.loc[wtm["OrdRep_WTM"] == "OrdRepetida_IMM", "F_TRABAJO_SESAP"] = pd.NaT

    # --- EstadoTrabajoFin_WTM ---
    tabla_wtm_estado = (
        wtm.drop_duplicates(subset=["NumeroOrdenWTM"], keep="first")
           .set_index("NumeroOrdenWTM")["EstadoTrabajo_WTM"]
    )
    encontrado = sesap["NumeroOrden"].map(tabla_wtm_estado)
    respaldo   = sesap["Cod_TipoIdent"].apply(lambda v: "PENDIENTE" if v == 5 else "SE")
    sesap["EstadoTrabajoFin_WTM"] = encontrado.where(encontrado.notna(), respaldo)

    # --- Estado_FHMAX y CT_FHMAX ---
    ahora = datetime.datetime.now()

    def calc_estado_fhmax(row):
        if str(row.get("ClaseDeOrden","")).strip() == "DC00":
            return ""
        if str(row.get("EstadoTrabajoFin_WTM","")).strip() != "PENDIENTE":
            return ""
        fhmax = row.get("FHMAX")
        if pd.isna(fhmax) if hasattr(pd,"isna") else fhmax is None:
            return ""
        try:
            fhmax_dt = pd.to_datetime(fhmax)
            return "ODS VENCIDA" if ahora > fhmax_dt else "PROX. VENCER"
        except Exception:
            return ""

    def calc_ct_fhmax(row):
        if str(row.get("ClaseDeOrden","")).strip() == "DC00":
            return ""
        if str(row.get("EstadoTrabajoFin_WTM","")).strip() != "PENDIENTE":
            return ""
        fhmax = row.get("FHMAX")
        if pd.isna(fhmax) if hasattr(pd,"isna") else fhmax is None:
            return ""
        try:
            fhmax_dt = pd.to_datetime(fhmax)
            delta = abs(ahora - fhmax_dt)
            horas_total = int(delta.total_seconds() // 3600)
            dias  = horas_total // 24
            horas = horas_total % 24
            return f"{dias:02d} {horas:02d}:00"
        except Exception:
            return ""

    sesap["Estado_FHMAX"] = sesap.apply(calc_estado_fhmax, axis=1)
    sesap["CT_FHMAX"]     = sesap.apply(calc_ct_fhmax,     axis=1)

    return wtm, sesap


def escribir_excel(wtm: pd.DataFrame, sesap: pd.DataFrame, ruta_salida: str):
    """Escribe BASES_PSR.xlsx con las hojas BD-WTM (tabla Excel) y BD-SESAP."""
    wb = Workbook(write_only=True)

    for nombre_hoja, df, columnas in (
        ("BD-WTM",   wtm,  COLUMNAS_WTM),
        ("BD-SESAP", sesap, COLUMNAS_SESAP),
    ):
        ws = wb.create_sheet(nombre_hoja)
        df = df.reindex(columns=columnas)

        n_filas = len(df) + 1   # +1 por la fila de encabezados
        n_cols  = len(columnas)

        # Encabezados
        ws.append(columnas)

        # Datos
        for fila in df.itertuples(index=False, name=None):
            ws.append([None if (v is None or (isinstance(v, float) and pd.isna(v))) else v
                       for v in fila])

        # Definir tabla Excel en BD-WTM (con TableColumn para modo write_only)
        if nombre_hoja == "BD-WTM":
            ref = f"A1:{get_column_letter(n_cols)}{n_filas}"
            tab = Table(displayName="BD_WTM", ref=ref)
            tab.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium9", showRowStripes=True
            )
            for i, nombre_col in enumerate(columnas, start=1):
                tab.tableColumns.append(TableColumn(id=i, name=nombre_col))
            ws.add_table(tab)

    os.makedirs(os.path.dirname(ruta_salida), exist_ok=True)
    wb.save(ruta_salida)
    print(f"\nArchivo guardado: {ruta_salida}")


# ------------------------------------------------------------------
# MAIN
# ------------------------------------------------------------------

def main():
    Path(CARPETA_BD_WTM).mkdir(parents=True, exist_ok=True)

    # --- 1. Descargar DC00, RC00, ZVCL en paralelo ---
    print("Descargando DC00, RC00 y ZVCL en paralelo...\n")
    errores = []
    with ThreadPoolExecutor(max_workers=3) as executor:
        futuros = {
            executor.submit(descargar_proceso, clave, datos): clave
            for clave, datos in PROCESOS.items()
        }
        for futuro in as_completed(futuros):
            clave = futuros[futuro]
            try:
                futuro.result()
                print(f">>> {clave} terminado con exito.")
            except Exception as e:
                print(f">>> ERROR en {clave}: {e}")
                errores.append(clave)

    if errores:
        print(f"\nAtencion: fallaron estos procesos: {errores}.")
        print("No se genera BASES_PSR.xlsx para evitar datos incompletos.")
        return

    # --- 2. Generar BASES_PSR.xlsx ---
    print("\nGenerando BASES_PSR.xlsx...")
    wtm   = construir_wtm()
    sesap = construir_sesap()
    wtm, sesap = aplicar_cruces(wtm, sesap)
    escribir_excel(wtm, sesap, RUTA_SALIDA)
    print("¡Listo!")


if __name__ == "__main__":
    main()
